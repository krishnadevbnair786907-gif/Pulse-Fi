import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

NAVY = RGBColor(18, 40, 76)
DARK = RGBColor(35, 35, 35)
ACCENT = RGBColor(0, 114, 178)


def add_slide(title_text, subtitle_speaker, bullets, img_path=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.3), Inches(0.9))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = NAVY

    p2 = tf.add_paragraph()
    p2.text = subtitle_speaker
    p2.font.size = Pt(13)
    p2.font.italic = True
    p2.font.color.rgb = ACCENT

    has_img = bool(img_path and os.path.exists(img_path))
    width_text = Inches(5.4) if has_img else Inches(12.0)
    bb = slide.shapes.add_textbox(Inches(0.5), Inches(1.35), width_text, Inches(5.6))
    btf = bb.text_frame
    btf.word_wrap = True

    for i, b in enumerate(bullets):
        bp = btf.paragraphs[0] if i == 0 else btf.add_paragraph()
        bp.text = f"•  {b}"
        bp.font.size = Pt(15)
        bp.font.color.rgb = DARK
        bp.space_after = Pt(12)

    if has_img:
        slide.shapes.add_picture(img_path, Inches(6.1), Inches(1.35), width=Inches(6.8))


slides_data = [
    (
        "PULSE-FI: Contactless Wi-Fi CSI Multi-Task Sensing",
        "Krishnadev B Nair (241140100) • Ardra Ajikumar (241140107) • Benert P Santosh (241140144) | Guide: Ms. Nikitha V",
        [
            "Device-free RF sensing using commodity ESP32 802.11n Channel State Information (CSI).",
            "Simultaneously solves three indoor sensing tasks without cameras or wearables:",
            "1. Presence & Activity Detection (Empty vs. Stationary vs. Moving): 100.00% Macro F1.",
            "2. Spatial RF Zone Fingerprinting (1m vs. 2m vs. 3m): 100.00% Macro F1.",
            "3. Respiratory Anomaly & BPM Monitoring (Normal vs. Fast/Tachypnea vs. Apnea): 100.00% F1.",
            "Validated on 12 hardware captures (12,492 CSI packets, 828 total evaluation windows)."
        ],
        "docs/figures/fig1_presence_waveforms.png"
    ),
    (
        "1. Problem Statement, Existing Systems & Objectives",
        "Speaker: Benert P Santosh & Ardra Ajikumar (Report Sections 1–4)",
        [
            "Existing Systems: Contact wearables (chest straps, pulse oximeters) cause discomfort; cameras raise severe privacy issues.",
            "Legacy RSSI Limitations: Single coarse scalar per packet suffers severe multipath fading and fails on micro-motion chest displacement.",
            "Proposed Pulse-Fi System: Extracts fine-grained amplitude across 52 active OFDM subcarriers.",
            "Core Objective: Build a leak-free signal processing & ensemble ML pipeline capable of detecting sub-centimeter chest motion and room occupancy in real time."
        ],
        "docs/figures/fig2_rf_zone_profiles.png"
    ),
    (
        "2. Hardware Architecture & CSI Data Acquisition (Days 1–4)",
        "Speaker: Krishnadev B Nair (Hardware & Pipeline Lead)",
        [
            "Transmitter / Receiver Setup: Dual ESP32 802.11n nodes operating in 2.4 GHz HT20 OFDM mode.",
            "Subcarrier Extraction: 64 total subcarriers parsed -> 12 null/guard subcarriers removed -> 52 active data subcarriers retained.",
            "Dataset Captured: 12 distinct hardware recordings totaling 12,492 raw CSI frames.",
            "Task Breakdown: 4 Presence recordings (312 windows), 3 Spatial Zone recordings (149 windows), and 5 Respiratory recordings (367 windows)."
        ],
        "docs/figures/fig3_respiratory_time_freq.png"
    ),
    (
        "3. Signal Conditioning & Subcarrier Selection (Days 1–4 / D1–D4)",
        "Speaker: Ardra Ajikumar (Signal Conditioning & Sequence Track)",
        [
            "D1 Uniform Resampling: Raw ESP32 UDP arrival averaged 12.37 Hz; cubic-spline interpolated onto a strict 20.0 Hz (dt = 50 ms) grid.",
            "D2 Hampel & Bandpass Denoising: Hampel outlier rejection + 4th-order zero-phase Butterworth filter (0.14–0.65 Hz) yielded +16.07 dB mean SNR gain.",
            "D3 Phase-Cancellation Proof: Selecting the Top-6 dynamic subcarriers preserves 1.25x–1.35x higher respiratory RMS than naive 52-subcarrier averaging.",
            "D4 Sliding Sub-Windows: 10 overlapping 3s steps (1s stride) per 12s window form the 3D tensor (367 x 10 x 17)."
        ],
        "docs/figures/fig9_ardra_d1_d4_pipeline.png"
    ),
    (
        "4. Leak-Free Physiological Feature Engineering & BPM Calibration",
        "Speaker: Krishnadev B Nair & Benert P Santosh",
        [
            "Session-Bias Audit: Removed static subcarrier DC ratios (sc_spatial_ratio, sc_spatial_std) from training to prevent posture/session memorization.",
            "17 Grounded Physiological Scalars: est_bpm, dom_freq, sig_rms, fast_rms, dropout_ratio, iqr, ptp, spectral_entropy, and band power fractions.",
            "112 Zero-Mean L2-Normalized AC Features: Captures relative spatial-dynamic fluctuation and velocity across all 52 subcarriers.",
            "Calibrated BPM Readout: Accurately separates Normal resting breathing (15.6 ± 3.2 BPM) from Fast/Tachypnea (22.1 ± 5.4 BPM) and Apnea (<5 BPM)."
        ],
        "docs/figures/fig4_bpm_distribution.png"
    ),
    (
        "5. PCA Feature Clustering & Multi-Task Separability",
        "Speaker: Krishnadev B Nair",
        [
            "Unsupervised 2D Principal Component Analysis (PCA) confirms strong class cluster separation prior to supervised classification.",
            "Presence & Activity: Empty room forms a tight noise-floor cluster, clearly separated from stationary breathing and high-variance walking.",
            "Spatial RF Zones: 1m, 2m, and 3m zones form distinct multipath comb manifolds.",
            "Respiratory States: Normal, Tachypnea, and Apnea occupy distinct regions in the 129-dimensional leak-free feature space."
        ],
        "docs/figures/fig5_pca_feature_clusters.png"
    ),
    (
        "6. Temporal Sequence / LSTM Modeling & Horizon Scaling (Days 5–8)",
        "Speaker: Ardra Ajikumar (Sequence & LSTM Lead)",
        [
            "Presence Sequence Tuning (D6–D8): Tuning window length (2s -> 4s), hidden layers (128, 64, 32), and L2 regularization improved Presence F1 from 81.53% to 100.00%.",
            "Respiratory Horizon Scaling (D5): Single 3s snapshot (T=1) achieves only 49.32% Macro F1 because 3s is shorter than one full breath cycle.",
            "10-Step Trajectory Gain: Modeling all 10 steps (T=10, 12s) with 17 features jumps F1 by +35.11% (to 84.43%).",
            "Bi-Recurrent Gate + Full 129-Feature Tensor: Achieves 100.00% 5-Fold CV Macro F1 across all 367 sequences."
        ],
        "docs/figures/fig10_d5_sequence_analysis.png"
    ),
    (
        "7. Random Forest Classifier & 3-Stage Feature Ablation (Days 6–8 / D8)",
        "Speaker: Benert P Santosh (Random Forest & Evaluation Lead)",
        [
            "Config A (17 Physiological Scalars Only): 85.29% CV Accuracy | 83.25% Macro F1 | 86.92% OOB Score.",
            "Config B (112 Zero-Mean AC Dynamic Profile): 99.73% CV Accuracy | 99.75% Macro F1 | 99.46% OOB Score.",
            "Config C (Combined 129 Leak-Free Set + GridSearchCV n_estimators=100): 100.00% 5-Fold Macro F1 and 100.00% Out-of-Bag accuracy.",
            "Top Gini Features: fast_to_sig_rms_ratio, log_p_fast_to_norm, p_fast_frac, est_bpm, and high-frequency AC velocity subcarriers."
        ],
        "docs/figures/fig8_d8_rf_analysis.png"
    ),
    (
        "8. Top-15 Gini Feature Importance Breakdown (Deliverable D8)",
        "Speaker: Benert P Santosh",
        [
            "Gini impurity analysis confirms the model relies on genuine physiological frequency and envelope dynamics rather than static room power.",
            "Frequency-Ratio Dominance: fast_to_sig_rms_ratio and log_p_fast_to_norm separate Tachypnea (>22 BPM) from resting breathing.",
            "Envelope & Dropout Features: min_sub_std, ptp, and dropout_ratio isolate Apnea breath-holds from low-amplitude resting respiration.",
            "Zero-Mean AC Subcarriers: Dynamic velocity on upper subcarriers (SC40–SC44) provides resilience against single-subcarrier multipath nulls."
        ],
        "docs/figures/fig7_rf_feature_importance.png"
    ),
    (
        "9. Unified Ensemble & Day 10 Error Analysis (Days 9–10 Milestone)",
        "Speaker: Ardra Ajikumar & Benert P Santosh",
        [
            "Day 9 Ensemble: Soft-voting combination of the Sequence Branch + Random Forest Branch achieves 99.75%–100.00% across all three tasks.",
            "Day 10 Error Analysis (59 Scalar Baseline Errors Audited):",
            "1. Apnea <-> Fast (26 errors): Noise-floor zero-crossing spikes during breath-hold mimic rapid shallow cycles in scalar FFT.",
            "2. Apnea <-> Normal (18 errors): Residual micro-motion at start/end of breath-hold mimics shallow resting breathing.",
            "3. Normal <-> Fast (15 errors): 2nd respiratory harmonic spillover; completely eliminated by the 129-feature ensemble."
        ],
        "docs/figures/fig11_d9_d10_ensemble_and_error_analysis.png"
    ),
    (
        "10. Multi-Task Confusion Matrices & Overall Benchmark Summary",
        "Speaker: Benert P Santosh & Krishnadev B Nair",
        [
            "Task 1 — Presence & Activity (312 windows): 100.00% Precision, Recall, and Macro F1 across Empty, Stationary, and Moving.",
            "Task 2 — Spatial RF Fingerprinting (149 windows): 100.00% Precision, Recall, and Macro F1 across Zone 1 (1m), Zone 2 (2m), and Zone 3 (3m).",
            "Task 3 — Respiratory Anomaly (367 windows): 100.00% Tuned RF / Bi-Recurrent Sequence F1 (171 Normal, 85 Fast, 111 Apnea).",
            "Zero data leakage verified via strict session-bias isolation and 5-fold stratified cross-validation."
        ],
        "docs/figures/fig6_confusion_matrices.png"
    ),
    (
        "11. Live Interactive Streamlit Dashboard & Conclusion (Days 11–14)",
        "Speaker: Krishnadev B Nair (Live Demo Walkthrough)",
        [
            "Real-Time Control Center (src/dashboard.py): Built in Streamlit + Plotly for live window scrubbing and multi-model inference.",
            "Simultaneous KPI Banner: Displays live Presence state, Spatial Zone (1m/2m/3m), Respiratory class, and calibrated BPM readout.",
            "4 Interactive Visual Panels: Top-3 respiratory waveforms, 52-subcarrier comb profile, Viridis spatio-temporal heatmap, and live Welch PSD spectrum.",
            "Future Scope: Multi-person blind source separation (ICA) and edge deployment via ESP-DL / TensorFlow Lite Micro."
        ],
        "docs/figures/fig3_respiratory_time_freq.png"
    )
]

for title, sub, b_list, img in slides_data:
    add_slide(title, sub, b_list, img)

os.makedirs("docs", exist_ok=True)
out_path = "docs/PulseFi_Presentation.pptx"
prs.save(out_path)
print(f"Successfully generated 12-slide presentation -> {out_path}")
